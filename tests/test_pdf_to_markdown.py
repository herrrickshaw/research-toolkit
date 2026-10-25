"""Tests for pdf_to_markdown.py using generated fixtures only.

No Zotero, no Dropbox, no real rclone: the end-to-end run() tests point RCLONE_BIN
at a stub script that copies between temp folders.
"""

import os
import stat
import sys
import textwrap
from pathlib import Path

import duckdb
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pdf_to_markdown as p2m  # noqa: E402


# ---- fixture builders -------------------------------------------------------

def make_pdf(path, text="Hello PDF world", pages=1):
    import fitz
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"{text} p{i + 1}")
    doc.save(str(path))
    doc.close()
    return path


def make_blank_pdf(path):
    import fitz
    doc = fitz.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()
    return path


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_paragraph("First paragraph")
    d.add_paragraph("   ")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "A", "B"
    d.save(str(path))
    return path


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Slide Title"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide -> skipped
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = s3.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "x", "y"
    prs.save(str(path))
    return path


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["a", None, 3])
    ws.append([None, None, None])
    wb.create_sheet("Empty")
    wb.save(str(path))
    return path


# ---- naming helpers -------------------------------------------------------------

@pytest.mark.parametrize("title,expected", [
    ("Report.pdf", "Report"),
    ("Deck.PPTX", "Deck"),
    ("a/b:c?.docx", "a_b_c_"),
    ("keep.dots v1.2", "keep.dots v1.2"),
    (None, "untitled"),
    ("", "untitled"),
])
def test_safe_title_for(title, expected):
    assert p2m.safe_title_for(title) == expected


def test_safe_title_for_truncates_to_80():
    assert len(p2m.safe_title_for("x" * 200)) == 80


@pytest.mark.parametrize("path,expected", [
    ("folder/file.pdf", ".pdf"),
    ("folder/file.tar.gz", ".gz"),
    ("folder/noext", ".bin"),
    ("folder/weird.p d!f", ".pdf"),
])
def test_ext_for(path, expected):
    assert p2m.ext_for(path) == expected


# ---- extractors ---------------------------------------------------------------

def test_extract_pdf(tmp_path):
    text, err = p2m.extract_text(make_pdf(tmp_path / "a.pdf", pages=2), "PDF")
    assert err is None
    assert "Hello PDF world p1" in text and "Hello PDF world p2" in text


def test_extract_blank_pdf_is_empty_not_error(tmp_path):
    text, err = p2m.extract_text(make_blank_pdf(tmp_path / "b.pdf"), "pdf")
    assert err is None and text.strip() == ""


def test_extract_docx_paragraphs_and_tables(tmp_path):
    text, err = p2m.extract_text(make_docx(tmp_path / "a.docx"), "docx")
    assert err is None
    assert text.split("\n\n") == ["First paragraph", "A | B"]


def test_extract_pptx_slides_and_tables(tmp_path):
    text, err = p2m.extract_text(make_pptx(tmp_path / "a.pptx"), "pptx")
    assert err is None
    assert "## Slide 1\nSlide Title" in text
    assert "## Slide 2" not in text          # empty slide skipped
    assert "## Slide 3\nx | y" in text


def test_extract_xlsx_rows_and_skips_empty(tmp_path):
    text, err = p2m.extract_text(make_xlsx(tmp_path / "a.xlsx"), "xlsx")
    assert err is None
    assert text == "## Sheet: Data\na |  | 3"


def test_corrupt_file_reports_error_not_exception(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    text, err = p2m.extract_text(bad, "docx")
    assert text is None and err.startswith("docx extraction error")


def test_unknown_type_uses_markitdown_fallback(tmp_path, monkeypatch):
    stub = tmp_path / "markitdown"
    stub.write_text("#!/bin/sh\necho converted:$1\n")
    stub.chmod(stub.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setattr(p2m, "MARKITDOWN_BIN", str(stub))
    f = tmp_path / "x.mobi"
    f.write_text("z")
    text, err = p2m.extract_text(f, "mobi")
    assert err is None and text.strip() == f"converted:{f}"


def test_markitdown_failure_is_error(tmp_path, monkeypatch):
    stub = tmp_path / "markitdown"
    stub.write_text("#!/bin/sh\necho boom >&2\nexit 3\n")
    stub.chmod(stub.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setattr(p2m, "MARKITDOWN_BIN", str(stub))
    text, err = p2m.extract_text(tmp_path, "html")
    assert text is None and err == "boom"


# ---- soffice batching -----------------------------------------------------------

def test_soffice_batches_one_call_per_filter_group(tmp_path, monkeypatch):
    calls = tmp_path / "calls.txt"
    stub = tmp_path / "soffice"
    # Fake soffice: log args, then write <stem>.<ext> for each input except ones named *fail*.
    stub.write_text(textwrap.dedent(f"""\
        #!/bin/sh
        echo "$@" >> {calls}
        filt=""; out=""
        while [ $# -gt 0 ]; do
          case "$1" in
            --convert-to) filt="$2"; shift 2;;
            --outdir) out="$2"; shift 2;;
            --*|-env:*) shift;;
            *) f="$1"; b=$(basename "$f"); s="${{b%.*}}"
               case "$filt" in csv*) e=csv;; *) e=txt;; esac
               case "$s" in *fail*) ;; *) echo "text of $s" > "$out/$s.$e";; esac
               shift;;
          esac
        done
        """))
    stub.chmod(stub.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setattr(p2m, "SOFFICE_BIN", str(stub))
    monkeypatch.setattr(p2m, "STAGE_DIR", tmp_path / "stage")
    items = []
    for key, ext, ft in [("k1", ".doc", "doc"), ("k2", ".ppt", "ppt"),
                         ("k3", ".xls", "xls"), ("kfail", ".doc", "doc")]:
        f = tmp_path / f"src_{key}{ext}"
        f.write_text("bin")
        items.append((key, f, ext, ft))
    res = p2m.convert_batch_soffice(items)
    assert len(calls.read_text().splitlines()) == 2      # doc+ppt together, xls separately
    assert res["k1"] == ("text of k1\n", None)
    assert res["k3"] == ("text of k3\n", None)
    assert res["kfail"] == (None, "soffice produced no output for this file")
    assert not (tmp_path / "stage" / "soffice_in").exists()


def test_soffice_empty_items():
    assert p2m.convert_batch_soffice([]) == {}


# ---- end-to-end run() with a stub rclone -----------------------------------------

@pytest.fixture
def pipeline(tmp_path, monkeypatch):
    cloud = tmp_path / "cloud"                 # pretend 'dropbox:' remote
    (cloud / "papers").mkdir(parents=True)
    uploaded = tmp_path / "uploaded"
    stub = tmp_path / "rclone"
    stub.write_text(textwrap.dedent(f"""\
        #!/bin/sh
        # rclone copy SRC DST [--files-from LIST] ...
        src="$2"; dst="$3"; list=""
        prev=""
        for a in "$@"; do [ "$prev" = "--files-from" ] && list="$a"; prev="$a"; done
        if [ -n "$list" ]; then
          while IFS= read -r p || [ -n "$p" ]; do
            [ -f "{cloud}/$p" ] || continue
            mkdir -p "$dst/$(dirname "$p")"; cp "{cloud}/$p" "$dst/$p"
          done < "$list"
        else
          mkdir -p "{uploaded}"; cp -R "$src"/. "{uploaded}/"
        fi
        """))
    stub.chmod(stub.stat().st_mode | stat.S_IEXEC)
    repo = tmp_path / "repo"
    repo.mkdir()
    monkeypatch.setattr(p2m, "RCLONE_BIN", str(stub))
    monkeypatch.setattr(p2m, "REPO_DIR", repo)
    monkeypatch.setattr(p2m, "DB_PATH", repo / "catalog.duckdb")
    monkeypatch.setattr(p2m, "MD_OUTPUT_DIR", repo / "md")
    monkeypatch.setattr(p2m, "STAGE_DIR", repo / "stage")
    monkeypatch.setattr(p2m, "FETCH_DIR", repo / "stage" / "fetch")
    monkeypatch.setattr(p2m, "UPLOAD_DIR", repo / "stage" / "upload")
    monkeypatch.setattr(p2m, "DISK_SAFETY_MARGIN_GB", 0)

    con = duckdb.connect(str(repo / "catalog.duckdb"))
    con.execute("create table documents (item_key varchar, md5 varchar, title varchar, path varchar,"
                " source varchar, filetype varchar, is_duplicate boolean)")
    con.close()

    def add(key, title, path, source="dropbox", filetype="pdf", dup=False, make=make_pdf):
        if make and source == "dropbox" and "!" not in path:
            make(cloud / path)
        c = duckdb.connect(str(repo / "catalog.duckdb"))
        c.execute("insert into documents values (?,?,?,?,?,?,?)",
                  [key, None, title, path, source, filetype, dup])
        c.close()

    def rows():
        c = duckdb.connect(str(repo / "catalog.duckdb"))
        out = {r[0]: r[1:] for r in c.execute(
            "select item_key, status, cloud_status, chars from fulltext_md").fetchall()}
        c.close()
        return out

    return SimpleNS(add=add, rows=rows, cloud=cloud, uploaded=uploaded, repo=repo)


class SimpleNS:
    def __init__(self, **kw):
        self.__dict__.update(kw)


def test_run_classifies_every_pending_item(pipeline):
    pipeline.add("A1", "Good.pdf", "papers/good.pdf")
    pipeline.add("A2", "Blank.pdf", "papers/blank.pdf", make=make_blank_pdf)
    pipeline.add("A3", "Inside zip", "papers/x.zip!inner.pdf")
    pipeline.add("A4", "Elsewhere", "foo/bar.pdf", source="onedrive")
    pipeline.add("A5", "Gone", "papers/missing.pdf", make=None)
    pipeline.add("A6", "Dup", "papers/dup.pdf", dup=True)
    p2m.run(max_docs=100, max_mb=1000)
    r = pipeline.rows()
    assert r["A1"][0] == "ok" and r["A1"][1] == "ok" and r["A1"][2] > 0
    assert r["A2"][0] == "empty_extraction"
    assert r["A3"][0] == "archive_not_persisted"
    assert r["A4"][0] == "unsupported_source"
    assert r["A5"][0] == "fetch_failed"
    assert "A6" not in r
    md = list((pipeline.repo / "md").glob("A1_*.md"))
    assert len(md) == 1 and "Hello PDF world" in md[0].read_text()
    assert sorted(p.name for p in pipeline.uploaded.iterdir()) == ["A1_Good.md", "A1_Good.pdf"]


def test_run_is_resumable_and_respects_max_docs(pipeline):
    for i in range(3):
        pipeline.add(f"B{i}", f"Doc{i}", f"papers/d{i}.pdf")
    p2m.run(max_docs=2, max_mb=1000)
    assert sorted(pipeline.rows()) == ["B0", "B1"]
    p2m.run(max_docs=2, max_mb=1000)
    assert sorted(pipeline.rows()) == ["B0", "B1", "B2"]


def test_run_respects_max_mb(pipeline):
    """--max-mb / PDF2MD_MAX_MB_PER_RUN is advertised as a per-run read cap. Once the
    cap is reached, remaining items must be left pending for the next run rather than
    converted in this one."""
    for i in range(3):
        pipeline.add(f"C{i}", f"Doc{i}", f"papers/c{i}.pdf")
    size_mb = (pipeline.cloud / "papers" / "c0.pdf").stat().st_size / (1024 * 1024)
    p2m.run(max_docs=100, max_mb=size_mb * 1.5)   # room for one file, not two
    assert sorted(pipeline.rows()) == ["C0"]
    p2m.run(max_docs=100, max_mb=1000)
    assert sorted(pipeline.rows()) == ["C0", "C1", "C2"]


def test_failed_upload_marks_cloud_failed_then_heals(pipeline, monkeypatch):
    pipeline.add("D1", "Doc", "papers/d.pdf")
    real = p2m.bulk_upload
    monkeypatch.setattr(p2m, "bulk_upload", lambda sources: {s: False for s in sources})
    p2m.run(max_docs=10, max_mb=1000)
    assert pipeline.rows()["D1"][:2] == ("ok", "failed")
    monkeypatch.setattr(p2m, "bulk_upload", real)
    p2m.run(max_docs=10, max_mb=1000)
    assert pipeline.rows()["D1"][:2] == ("ok", "ok")
